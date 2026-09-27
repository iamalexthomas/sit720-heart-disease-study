"""Format the report text as an editable Word document.

This is called by build_report.py and does not run any experiments.
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Mm, Pt, RGBColor
from PIL import Image

REPORT_FOLDER = Path(__file__).resolve().parents[1] / "report"


def add_text(paragraph, text):
    """Keep web addresses clickable in Word as well as in the PDF."""
    for part in re.split(r"(https?://\S+)", text):
        if not part.startswith(("https://", "http://")):
            paragraph.add_run(part)
            continue
        url = part.rstrip(".")
        link = OxmlElement("w:hyperlink")
        relationship = paragraph.part.relate_to(
            url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        link.set(qn("r:id"), relationship)
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        colour = OxmlElement("w:color")
        colour.set(qn("w:val"), "000000")
        properties.append(colour)
        run.append(properties)
        value = OxmlElement("w:t")
        value.text = url
        run.append(value)
        link.append(run)
        paragraph._p.append(link)
        if part.endswith("."):
            paragraph.add_run(".")


def add_table(document, rows):
    count = len(rows[0])
    width = 178
    if count == 2:
        widths = [50, 128]
    elif count == 3:
        widths = [43, 70, 65]
    elif count == 7:
        widths = [19, 23] + [27.2] * 5
    else:
        widths = [24] + [(width - 24) / (count - 1)] * (count - 1)
    table = document.add_table(rows=0, cols=count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, column_width in zip(table.columns, widths):
        column.width = Mm(column_width)
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement("w:" + edge)
        for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
            border.set(qn("w:" + key), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for column_index, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Mm(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                margin = OxmlElement("w:" + side)
                margin.set(qn("w:w"), "70")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            cell_properties.append(margins)
            if index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E7EDF3")
                cell_properties.append(shading)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.keep_with_next = index == 0
            if count >= 3 and column_index > 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(value)
            run.font.size = Pt(8.5 if count >= 6 else 9.5)
            run.bold = index == 0
    return table


def render_docx(markdown, output):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(16)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(17)
    section.footer_distance = Mm(8)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.05
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 20), ("Heading 1", 13), ("Heading 2", 11.5)]:
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    caption = document.styles["Caption"]
    caption.font.name = "Arial"
    caption.font.size = Pt(9)
    caption.font.bold = False
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(64, 64, 64)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(7)
    document.core_properties.title = "SIT720 Machine Learning Mini Research"
    document.core_properties.author = "Alex Thomas"
    document.core_properties.subject = "Heart disease classification study"
    document.core_properties.comments = ""
    # The default template may have a coloured line under its title style.
    for style in document.styles:
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)

    footer = section.footer.paragraphs[0]
    footer.style = "Normal"
    footer.paragraph_format.tab_stops.add_tab_stop(Mm(178), WD_TAB_ALIGNMENT.RIGHT)
    footer.add_run("Alex Thomas | SIT720 11.1HD\t")
    page_number = OxmlElement("w:fldSimple")
    page_number.set(qn("w:instr"), "PAGE")
    footer._p.append(page_number)
    for run in footer.runs:
        run.font.size = Pt(8)

    lines = markdown.strip().splitlines()
    index = 0
    new_page = False
    while index < len(lines):
        line = lines[index].strip()
        if line == "<!-- pagebreak -->":
            new_page = True
        elif line.startswith("# "):
            document.add_paragraph(line[2:], "Title")
        elif line.startswith("## "):
            heading = line[3:]
            level = 2 if re.match(r"\d+\.\d+", heading) else 1
            paragraph = document.add_heading(heading, level)
            paragraph.paragraph_format.page_break_before = new_page
            new_page = False
        elif line.startswith("| "):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [value.strip() for value in lines[index].strip().strip("|").split("|")]
                if not all(value == "---" for value in cells):
                    rows.append(cells)
                index += 1
            add_table(document, rows)
            continue
        elif line.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            image_path = REPORT_FOLDER / match.group(2)
            with Image.open(image_path) as picture:
                image_width, image_height = picture.size
            height_limit = 69 if "roc_curves" not in line else 73
            width = min(178, height_limit * image_width / image_height)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_after = Pt(0)
            shape = paragraph.add_run().add_picture(str(image_path), width=Mm(width))
            shape._inline.docPr.set("descr", match.group(1))
            document.add_paragraph(match.group(1), "Caption")
        elif line:
            style = "Caption" if line.startswith("Table ") else "Normal"
            paragraph = document.add_paragraph(style=style)
            add_text(paragraph, line)
        index += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
